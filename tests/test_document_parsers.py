import asyncio
import io
import logging

import pytest
from docx import Document
from ebooklib import epub

from app.core.exceptions import DocumentOutlineRequiredError
from app.core.logger import _redact_message
from app.services.document_parser_service import extract_chapters_by_document_type, parse_document_structure
from app.services.document_structure_types import flatten_sections
from app.services.docx_service import extract_docx_chapters, parse_docx_structure
from app.services.epub_service import extract_epub_chapters, parse_epub_structure


def _run(coro):
    return asyncio.run(coro)


def _build_docx_with_headings() -> io.BytesIO:
    document = Document()
    document.add_heading("Chương 1: Mở đầu", level=1)
    document.add_paragraph("Nội dung chương một.")
    document.add_heading("Chương 2: Kết luận", level=1)
    document.add_paragraph("Nội dung chương hai.")

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


def _build_nested_docx() -> io.BytesIO:
    document = Document()
    document.add_heading("Chương 1", level=1)
    document.add_paragraph("Mở đầu chương.")
    document.add_heading("Mục 1.1", level=2)
    document.add_paragraph("Nội dung mục một một.")
    document.add_heading("Tiểu mục 1.1.1", level=3)
    document.add_paragraph("Nội dung tiểu mục.")

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


def _build_docx_without_headings() -> io.BytesIO:
    document = Document()
    document.add_paragraph("Đây chỉ là đoạn văn thường, không có heading.")

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


def _build_epub(tmp_path, with_toc: bool) -> io.BytesIO:
    book = epub.EpubBook()
    book.set_identifier("triam-test")
    book.set_title("TRIAM Test EPUB")
    book.set_language("vi")

    chapter = epub.EpubHtml(title="Chương 1: Mở đầu", file_name="chap_01.xhtml", lang="vi")
    chapter.content = "<h1>Chương 1: Mở đầu</h1><p>Nội dung chương một.</p>"
    book.add_item(chapter)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", chapter]

    if with_toc:
        book.toc = (epub.Link("chap_01.xhtml", "Chương 1: Mở đầu", "chap_01"),)

    epub_path = tmp_path / ("with_toc.epub" if with_toc else "without_toc.epub")
    epub.write_epub(str(epub_path), book)
    return io.BytesIO(epub_path.read_bytes())


def _build_nested_epub(tmp_path) -> io.BytesIO:
    book = epub.EpubBook()
    book.set_identifier("triam-nested-test")
    book.set_title("TRIAM Nested EPUB")
    book.set_language("vi")

    chapter = epub.EpubHtml(title="Chương 1", file_name="chap_01.xhtml", lang="vi")
    chapter.content = "<h1>Chương 1</h1><h2>Mục 1.1</h2><p>Nội dung mục một.</p>"
    child = epub.EpubHtml(title="Mục 1.1", file_name="chap_01_01.xhtml", lang="vi")
    child.content = "<h2>Mục 1.1</h2><p>Nội dung mục con.</p>"
    book.add_item(chapter)
    book.add_item(child)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", chapter, child]
    book.toc = ((epub.Link("chap_01.xhtml", "Chương 1", "chap_01"), (epub.Link("chap_01_01.xhtml", "Mục 1.1", "chap_01_01"),)),)

    epub_path = tmp_path / "nested.epub"
    epub.write_epub(str(epub_path), book)
    return io.BytesIO(epub_path.read_bytes())


def test_extract_pdf_chapters_from_sample(sample_pdf):
    with open(sample_pdf, "rb") as file_obj:
        chapters = _run(extract_chapters_by_document_type(file_obj, "pdf"))

    assert len(chapters) == 2
    assert chapters[0]["title"] == "Chapter 1: The Science of Memory"
    assert chapters[0]["page_start"] == 1
    assert chapters[0]["text"]


def test_extract_docx_chapters_with_headings():
    chapters = _run(extract_docx_chapters(_build_docx_with_headings()))

    assert [chapter["title"] for chapter in chapters] == ["Chương 1: Mở đầu", "Chương 2: Kết luận"]
    assert chapters[0]["page_start"] is None
    assert "Nội dung chương một" in chapters[0]["text"]


def test_parse_docx_structure_preserves_heading_tree():
    parsed = _run(parse_docx_structure(_build_nested_docx()))
    flattened = flatten_sections(parsed.sections)

    assert parsed.total_sections == 3
    assert parsed.max_depth == 3
    assert [(section.path, section.level, section.title) for section in flattened] == [
        ("1", 1, "Chương 1"),
        ("1.1", 2, "Mục 1.1"),
        ("1.1.1", 3, "Tiểu mục 1.1.1"),
    ]


def test_extract_docx_without_headings_fails():
    with pytest.raises(DocumentOutlineRequiredError):
        _run(extract_docx_chapters(_build_docx_without_headings()))


def test_extract_epub_chapters_with_toc(tmp_path):
    chapters = _run(extract_epub_chapters(_build_epub(tmp_path, with_toc=True)))

    assert [chapter["title"] for chapter in chapters] == ["Chương 1: Mở đầu"]
    assert chapters[0]["page_start"] is None
    assert "Nội dung chương một" in chapters[0]["text"]


def test_parse_epub_structure_preserves_nested_toc(tmp_path):
    parsed = _run(parse_epub_structure(_build_nested_epub(tmp_path)))
    flattened = flatten_sections(parsed.sections)

    assert parsed.total_sections == 2
    assert parsed.max_depth == 2
    assert [(section.path, section.level, section.title) for section in flattened] == [
        ("1", 1, "Chương 1"),
        ("1.1", 2, "Mục 1.1"),
    ]


def test_parse_document_structure_dispatches_pdf(sample_pdf):
    with open(sample_pdf, "rb") as file_obj:
        parsed = _run(parse_document_structure(file_obj, "pdf"))

    assert parsed.document_type == "pdf"
    assert parsed.total_sections == 2
    assert parsed.sections[0].title == "Chapter 1: The Science of Memory"


def test_extract_epub_without_toc_fails(tmp_path):
    with pytest.raises(DocumentOutlineRequiredError):
        _run(extract_epub_chapters(_build_epub(tmp_path, with_toc=False)))


def test_redact_message_hides_sensitive_headers():
    message = (
        "{'authorization': 'Bearer eyJ.secret.token', 'apikey': 'supabase-secret', "
        "'api_key': 'livekit-key', 'api_secret': 'livekit-secret', "
        "'access_token': 'access-secret', 'refresh_token': 'refresh-secret'}"
    )

    redacted = _redact_message(message)

    for secret in (
        "eyJ.secret.token",
        "supabase-secret",
        "livekit-key",
        "livekit-secret",
        "access-secret",
        "refresh-secret",
    ):
        assert secret not in redacted
    assert "[REDACTED]" in redacted


def test_noisy_http_loggers_are_warning_or_higher():
    assert logging.getLogger("httpcore").level >= logging.WARNING
    assert logging.getLogger("httpx").level >= logging.WARNING
    assert logging.getLogger("hpack").level >= logging.WARNING
