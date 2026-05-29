from __future__ import annotations

from dataclasses import dataclass, field

from docx import Document
from docx.oxml.ns import qn

from app.core.exceptions import DocumentExtractionError, DocumentOutlineRequiredError
from app.services.document_structure_types import ParsedDocument, ParsedSection, parsed_sections_to_legacy_chapters

HEADING_STYLE_NAMES = {
    "heading 1",
    "heading 2",
    "heading 3",
}


@dataclass
class _SectionBuilder:
    title: str
    level: int
    order_index: int
    path: str
    parts: list[str] = field(default_factory=list)
    children: list["_SectionBuilder"] = field(default_factory=list)

    def to_parsed(self) -> ParsedSection:
        return ParsedSection(
            title=self.title,
            level=self.level,
            order_index=self.order_index,
            path=self.path,
            text="\n".join(self.parts).strip(),
            page_start=None,
            page_end=None,
            children=[child.to_parsed() for child in self.children],
        )


def _has_outline_level(paragraph) -> bool:
    if paragraph.style is None or paragraph.style.element is None:
        return False

    p_pr = paragraph.style.element.find(qn("w:pPr"))
    if p_pr is None:
        return False

    return p_pr.find(qn("w:outlineLvl")) is not None


def _outline_level(paragraph) -> int | None:
    if paragraph.style is None or paragraph.style.element is None:
        return None

    p_pr = paragraph.style.element.find(qn("w:pPr"))
    if p_pr is None:
        return None

    outline = p_pr.find(qn("w:outlineLvl"))
    if outline is None:
        return None

    value = outline.get(qn("w:val"))
    if value is None:
        return None

    try:
        return int(value) + 1
    except ValueError:
        return None


def _heading_level(paragraph) -> int | None:
    style_name = (paragraph.style.name or "").strip().lower() if paragraph.style else ""
    if style_name.startswith("heading "):
        suffix = style_name.removeprefix("heading ").strip()
        if suffix.isdigit():
            return max(1, min(6, int(suffix)))

    outline_level = _outline_level(paragraph)
    if outline_level is not None:
        return max(1, min(6, outline_level))

    return None


def _is_heading(paragraph) -> bool:
    style_name = (paragraph.style.name or "").strip().lower() if paragraph.style else ""
    return style_name in HEADING_STYLE_NAMES or style_name.startswith("heading ") or _has_outline_level(paragraph)


async def parse_docx_structure(file_obj) -> ParsedDocument:
    file_obj.seek(0)

    try:
        document = Document(file_obj)
    except Exception as exc:
        raise DocumentExtractionError(f"Không thể đọc DOCX: {exc}") from exc

    roots: list[_SectionBuilder] = []
    stack: list[_SectionBuilder] = []
    sibling_counts: dict[str, int] = {"": 0}

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue

        if _is_heading(paragraph):
            level = _heading_level(paragraph) or 1
            while stack and stack[-1].level >= level:
                stack.pop()

            parent = stack[-1] if stack else None
            parent_path = parent.path if parent else ""
            sibling_counts[parent_path] = sibling_counts.get(parent_path, 0) + 1
            order_index = sibling_counts[parent_path] - 1
            path = f"{parent_path}.{order_index + 1}" if parent_path else str(order_index + 1)
            builder = _SectionBuilder(title=text, level=level, order_index=order_index, path=path)

            if parent:
                parent.children.append(builder)
            else:
                roots.append(builder)

            stack.append(builder)
            continue

        if stack:
            stack[-1].parts.append(text)

    if not roots:
        raise DocumentOutlineRequiredError("DOCX không có Heading/mục lục hợp lệ. Vui lòng dùng Heading 1/Heading 2 cho tên chương.")

    return ParsedDocument(document_type="docx", sections=[root.to_parsed() for root in roots])


async def extract_docx_chapters(file_obj) -> list[dict]:
    parsed = await parse_docx_structure(file_obj)
    chapters = parsed_sections_to_legacy_chapters(parsed.sections)
    if not chapters:
        raise DocumentExtractionError("DOCX có Heading nhưng không trích xuất được nội dung text.")
    return chapters
